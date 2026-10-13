"""hlc_eod_report.py — End-of-day report generator (PDF + XLSX + CSV + JSON)"""
import io, json, logging
from datetime import datetime
from typing import List, Dict, Any
logger = logging.getLogger(__name__)

def assemble_eod_data(team_db, visitor_db, report_date: str) -> Dict:
    return {
        "report_date":      report_date,
        "generated_at":     datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "team_stats":       team_db.get_stats(report_date),
        "team_attendance":  team_db.get_attendance(report_date),
        "visitor_stats":    visitor_db.get_stats(report_date),
        "visitor_sessions": visitor_db.get_sessions(visit_date=report_date),
        "visitor_summary":  visitor_db.get_daily_summary(report_date),
        "over_15_list":     [s for s in visitor_db.get_sessions(visit_date=report_date) if s.get("over_15min")],
        "alerts":           visitor_db.get_all_alerts(visit_date=report_date),
        "open_sessions":    visitor_db.get_all_open_sessions(report_date),
    }

def generate_eod_json(data: Dict) -> bytes:
    return json.dumps(data,ensure_ascii=False,indent=2,default=str).encode("utf-8")

def generate_eod_csv_zip(data: Dict) -> bytes:
    import zipfile, io as _io, csv
    def _csv(rows):
        if not rows: return ""
        buf=_io.StringIO()
        w=csv.DictWriter(buf,fieldnames=list(rows[0].keys()),extrasaction="ignore")
        w.writeheader(); w.writerows(rows)
        return buf.getvalue()
    buf=_io.BytesIO()
    with zipfile.ZipFile(buf,"w",zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("hlc_team_attendance.csv",  _csv(data["team_attendance"]))
        zf.writestr("hlc_visitor_sessions.csv",  _csv(data["visitor_sessions"]))
        zf.writestr("hlc_visitor_summary.csv",   _csv(data["visitor_summary"]))
        zf.writestr("hlc_alerts.csv",            _csv(data["alerts"]))
    return buf.getvalue()

def generate_eod_xlsx(data: Dict) -> bytes:
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        raise RuntimeError("Run: pip install openpyxl --break-system-packages")
    wb=openpyxl.Workbook(); rdt=data["report_date"]
    NAVY,GOLD,WHITE="1A3A6B","C8A84B","FFFFFF"

    def _hdr(ws,row,col,val):
        c=ws.cell(row=row,column=col,value=val)
        c.font=Font(name="Calibri",bold=True,color=WHITE,size=9)
        c.fill=PatternFill("solid",fgColor=NAVY)
        c.alignment=Alignment(horizontal="center",vertical="center")

    def _cell(ws,row,col,val):
        c=ws.cell(row=row,column=col,value=val)
        c.font=Font(name="Calibri",size=9)

    # Sheet 1: Team Attendance
    ws1=wb.active; ws1.title="Team Attendance"
    ws1.merge_cells("A1:H1")
    t=ws1.cell(row=1,column=1,value=f"HLC Team Attendance — {rdt}")
    t.font=Font(name="Calibri",bold=True,color=GOLD,size=12)
    t.fill=PatternFill("solid",fgColor=NAVY)
    hdrs=["ID","Name","Dept","Role","Date","First In","Last Out","Status"]
    for ci,h in enumerate(hdrs,1): _hdr(ws1,2,ci,h)
    for ri,r in enumerate(data["team_attendance"],3):
        for ci,k in enumerate(["employee_id","name","dept","hlc_role","work_date","first_seen","last_seen","final_status"],1):
            _cell(ws1,ri,ci,r.get(k,""))

    # Sheet 2: Visitor Sessions
    ws2=wb.create_sheet("Visitor Log")
    ws2.merge_cells("A1:J1")
    t2=ws2.cell(row=1,column=1,value=f"HLC Visitor Log — {rdt}")
    t2.font=Font(name="Calibri",bold=True,color=GOLD,size=12)
    t2.fill=PatternFill("solid",fgColor=NAVY)
    hdrs2=["ID","Visitor ID","Name","Dept","Entry","Exit","Duration(min)","Visit#","Over15?","Door"]
    for ci,h in enumerate(hdrs2,1): _hdr(ws2,2,ci,h)
    for ri,r in enumerate(data["visitor_sessions"],3):
        for ci,k in enumerate(["id","visitor_id","name","dept","entry_time","exit_time","duration_minutes","visit_number","over_15min","door_id"],1):
            _cell(ws2,ri,ci,r.get(k,""))

    buf=io.BytesIO(); wb.save(buf); return buf.getvalue()

def generate_eod_pdf(data: Dict) -> bytes:
    try:
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib import colors
        from reportlab.lib.units import cm
        from reportlab.platypus import SimpleDocTemplate,Table,TableStyle,Paragraph,Spacer,HRFlowable
        from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
        from reportlab.lib.enums import TA_CENTER
    except ImportError:
        raise RuntimeError("Run: pip install reportlab --break-system-packages")

    NAVY=colors.HexColor("#1A3A6B"); GOLD=colors.HexColor("#C8A84B")
    GREEN=colors.HexColor("#16A34A"); RED=colors.HexColor("#DC2626")
    LGREY=colors.HexColor("#F0F4F8"); WHITE=colors.white
    buf=io.BytesIO()
    doc=SimpleDocTemplate(buf,pagesize=A4,leftMargin=1.5*cm,rightMargin=1.5*cm,topMargin=2*cm,bottomMargin=1.5*cm)
    styles=getSampleStyleSheet()
    title_s=ParagraphStyle("T",parent=styles["Normal"],fontSize=14,fontName="Helvetica-Bold",textColor=WHITE,alignment=TA_CENTER)
    sec_s=ParagraphStyle("S",parent=styles["Normal"],fontSize=11,fontName="Helvetica-Bold",textColor=NAVY)
    foot_s=ParagraphStyle("F",parent=styles["Normal"],fontSize=7,textColor=colors.HexColor("#8A9AB0"),alignment=TA_CENTER)
    rdt=data["report_date"]; ts=data["team_stats"]; vs=data["visitor_stats"]
    story=[]

    # Header
    hdr=Table([[Paragraph(f"<font color='#C8A84B'>EGS Pillay Engineering College</font><br/><b>HLC End-of-Day Report — {rdt}</b><br/><font size=9 color='#C8D8F0'>Generated: {data['generated_at']}</font>",title_s)]],colWidths=[doc.width])
    hdr.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,-1),NAVY),("TOPPADDING",(0,0),(-1,-1),14),("BOTTOMPADDING",(0,0),(-1,-1),14)]))
    story.append(hdr); story.append(Spacer(1,0.4*cm))

    # Summary table
    story.append(Paragraph("Summary",sec_s)); story.append(Spacer(1,0.2*cm))
    summ=[["Metric","Value"],
          ["HLC Team Enrolled",ts.get("enrolled","—")],
          ["Present Today",ts.get("present",0)],
          ["Absent Today",ts.get("absent",0)],
          ["Unique Visitors",vs.get("unique_visitors",0)],
          ["Total Visitor Entries",vs.get("total_visits",0)],
          ["Visitors Over 15 min",vs.get("over_15_count",0)],
          ["Open Alerts",len(data["alerts"])],
          ["Generated At",data["generated_at"]]]
    st=Table(summ,colWidths=[5*cm,4*cm])
    st.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),NAVY),("TEXTCOLOR",(0,0),(-1,0),WHITE),
        ("FONTNAME",(0,0),(-1,0),"Helvetica-Bold"),("FONTSIZE",(0,0),(-1,-1),9),
        ("ROWBACKGROUNDS",(0,1),(-1,-1),[WHITE,LGREY]),("GRID",(0,0),(-1,-1),0.3,colors.HexColor("#DDE3EB")),
        ("TOPPADDING",(0,0),(-1,-1),4),("BOTTOMPADDING",(0,0),(-1,-1),4)]))
    story.append(st); story.append(Spacer(1,0.4*cm))

    # Team attendance
    if data["team_attendance"]:
        story.append(Paragraph("Team Attendance",sec_s)); story.append(Spacer(1,0.2*cm))
        ta=[["ID","Name","Dept","Role","First In","Last Out","Status"]]
        for r in data["team_attendance"]:
            ta.append([r.get("employee_id",""),r.get("name",""),r.get("dept",""),r.get("hlc_role",""),r.get("first_seen","—"),r.get("last_seen","—"),r.get("final_status","ABSENT")])
        ta_tbl=Table(ta,colWidths=[2*cm,4*cm,3*cm,3*cm,2.2*cm,2.2*cm,2.2*cm])
        ta_tbl.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),NAVY),("TEXTCOLOR",(0,0),(-1,0),WHITE),
            ("FONTNAME",(0,0),(-1,0),"Helvetica-Bold"),("FONTSIZE",(0,0),(-1,-1),8),
            ("ROWBACKGROUNDS",(0,1),(-1,-1),[WHITE,LGREY]),("GRID",(0,0),(-1,-1),0.3,colors.HexColor("#DDE3EB")),
            ("TOPPADDING",(0,0),(-1,-1),3),("BOTTOMPADDING",(0,0),(-1,-1),3)]))
        story.append(ta_tbl); story.append(Spacer(1,0.3*cm))

    # Footer
    story.append(Spacer(1,0.5*cm))
    story.append(HRFlowable(width="100%",thickness=0.5,color=GOLD))
    story.append(Paragraph(f"EGS Pillay Engineering College · Nagapattinam | FaceTrack Pro v2.0 | {data['generated_at']}",foot_s))
    doc.build(story); return buf.getvalue()
